"""
将 docs/毕业论文预演稿.md 转换为规范排版的 Word 文档
（标题层级、正文小四宋体、图注表注、参考文献悬挂缩进）
输出：docs/毕业论文预演稿.docx
"""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD = os.path.join(BASE, "docs", "毕业论文预演稿.md")
OUT = os.path.join(BASE, "docs", "毕业论文预演稿.docx")
FIGDIR = os.path.join(BASE, "figures")


def set_font(run, name="宋体", size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color


def body_para(doc, text, indent=True, size=12):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    if indent:
        pf.first_line_indent = Pt(size * 2)
    # 处理 **粗体**
    for seg in re.split(r"(\*\*[^*]+\*\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            set_font(p.add_run(seg[2:-2]), "宋体", size, bold=True)
        else:
            set_font(p.add_run(seg), "宋体", size)
    return p


def main():
    src = open(MD, encoding="utf-8").read().split("\n")
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.0)
    st = doc.styles["Normal"]
    st.font.name = "宋体"; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    i, n = 0, len(src)
    while i < n:
        line = src[i].rstrip()
        if not line.strip():
            i += 1
            continue
        # 图片
        m = re.match(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if m:
            rel = m.group(2)
            path = os.path.join(BASE, rel)
            if os.path.exists(path):
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=Cm(13.5))
                cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(cap.add_run(m.group(1)), "宋体", 10.5)
            i += 1
            continue
        # 表格
        if line.strip().startswith("|") and i + 1 < n and set(src[i + 1].replace("|", "").strip()) <= {"-", ":", " "}:
            rows = []
            while i < n and src[i].strip().startswith("|"):
                cells = [c.strip() for c in src[i].strip().strip("|").split("|")]
                if not set("".join(cells)) <= {"-", ":", " "}:
                    rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r, row in enumerate(rows):
                    for c, val in enumerate(row[:len(rows[0])]):
                        cell = t.cell(r, c)
                        cell.text = ""
                        para = cell.paragraphs[0]
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        set_font(para.add_run(re.sub(r"\*\*", "", val)),
                                 "宋体", 9.5, bold=(r == 0))
            doc.add_paragraph()
            continue
        # 引用块
        if line.strip().startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            set_font(p.add_run(re.sub(r"[>*]", "", line).strip()), "楷体", 11)
            i += 1
            continue
        # 标题
        if line.startswith("#### "):
            p = doc.add_paragraph(); set_font(p.add_run(line[5:]), "黑体", 12, bold=True)
        elif line.startswith("### "):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
            set_font(p.add_run(line[4:]), "黑体", 14, bold=True)
        elif line.startswith("## "):
            txt = line[3:]
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(8)
            set_font(p.add_run(txt), "黑体", 16, bold=True)
        elif line.startswith("# "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(line[2:]), "黑体", 18, bold=True)
        elif line.strip() == "---":
            i += 1; continue
        # 代码块 / 伪代码
        elif line.strip().startswith("```"):
            i += 1
            buf = []
            while i < n and not src[i].strip().startswith("```"):
                buf.append(src[i]); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.line_spacing = 1.0
            set_font(p.add_run("\n".join(buf)), "Consolas", 8.5)
            i += 1
            continue
        # 公式
        elif line.strip().startswith("$$") or line.strip().startswith("$"):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(line.strip().strip("$")), "Cambria Math", 11)
        # 参考文献条目
        elif re.match(r"^\[\d+\]", line.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            set_font(p.add_run(re.sub(r"<sup>.*?</sup>", "", line).strip()), "宋体", 10.5)
        elif line.strip().startswith("- ") or line.strip().startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            body_para(doc, re.sub(r"^[-*] ", "", line.strip()), indent=False)
            doc.paragraphs[-1]._element.getparent().remove(doc.paragraphs[-1]._element)
            i += 1
            continue
        else:
            body_para(doc, re.sub(r"<sup>(.*?)</sup>", r"\1", line.strip()))
        i += 1

    doc.save(OUT)
    print(f"[OK] 已生成 {OUT}")


if __name__ == "__main__":
    main()
